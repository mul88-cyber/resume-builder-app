import streamlit as st
import json
from fpdf import FPDF
import io
import zipfile
from datetime import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import base64
import tempfile
import os

# --- KONFIGURASI HALAMAN ---
st.set_page_config(
    page_title="CV Builder Pro Ultra v2.1",
    page_icon="🚀",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- STATE MANAGEMENT ---
if 'cv_data' not in st.session_state:
    st.session_state.cv_data = {
        'personal_info': {
            'nama': '', 'email': '', 'telepon': '', 'alamat': '',
            'linkedin': '', 'github': '', 'website': '', 'posisi_target': '',
            'foto': None # Base64 string
        },
        'ringkasan': '',
        'pengalaman': [],
        'pendidikan': [],
        'keahlian': [],
        'sertifikasi': [],
        'proyek': [],
        'bahasa': [],
        'hobi': []
    }

if 'settings' not in st.session_state:
    st.session_state.settings = {
        'template_style': 'modern_sidebar',
        'font_family': 'Helvetica',
        'base_color': '#2563eb',
        'accent_color': '#1e40af',
        'font_size_body': 10,
        'font_size_header': 24,
        'theme': 'light'
    }

# --- DATA TEMPLATES ---
LAYOUTS = {
    'modern_sidebar': {'name': 'Modern Sidebar', 'type': '2_column', 'ats_score': 85},
    'classic_vertical': {'name': 'ATS Professional', 'type': '1_column', 'ats_score': 95},
    'executive': {'name': 'Executive Style', 'type': '2_column', 'ats_score': 90},
    'creative': {'name': 'Creative Portfolio', 'type': 'creative', 'ats_score': 70}
}

FONTS = {
    'Helvetica': {'pdf': 'Helvetica', 'docx': 'Arial'},
    'Times': {'pdf': 'Times', 'docx': 'Times New Roman'},
    'Courier': {'pdf': 'Courier', 'docx': 'Courier New'}
}

# --- FUNGSI HELPER ---
def hex_to_rgb(hex_color):
    """Safe Hex to RGB converter"""
    try:
        hex_color = hex_color.lstrip('#')
        if len(hex_color) == 6:
            return tuple(int(hex_color[i:i+2], 16) for i in (0, 2, 4))
        return (37, 99, 235) # Default Blue
    except:
        return (37, 99, 235)

def save_image_temp(base64_str):
    """Decodes base64 image to a temp file for PDF/Docx inclusion"""
    if not base64_str:
        return None
    try:
        # Remove header if present (data:image/png;base64,...)
        if "," in base64_str:
            base64_str = base64_str.split(",")[1]
        
        image_data = base64.b64decode(base64_str)
        tfile = tempfile.NamedTemporaryFile(delete=False, suffix='.png')
        tfile.write(image_data)
        tfile.close()
        return tfile.name
    except Exception as e:
        print(f"Error saving image: {e}")
        return None

def calculate_ats_score(data):
    score = 0
    if data['personal_info']['nama']: score += 10
    if data['personal_info']['email']: score += 10
    if data['personal_info']['posisi_target']: score += 10
    if data['ringkasan']: score += 15
    if len(data['pengalaman']) >= 1: score += 20
    if len(data['pendidikan']) >= 1: score += 15
    if len(data['keahlian']) >= 5: score += 20
    return min(score, 100)

# --- PDF GENERATOR ---
class CVPDF(FPDF):
    def footer(self):
        self.set_y(-15)
        self.set_font('Helvetica', 'I', 8)
        self.set_text_color(128)
        self.cell(0, 10, f"Page {self.page_no()}", 0, 0, 'C')

def generate_pdf_enhanced(data, settings):
    pdf = CVPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    
    # Colors & Fonts
    r_prim, g_prim, b_prim = hex_to_rgb(settings['base_color'])
    r_sec, g_sec, b_sec = hex_to_rgb(settings['accent_color'])
    font = FONTS.get(settings['font_family'], FONTS['Helvetica'])['pdf']
    
    # Handle Photo (Temp File)
    temp_img_path = save_image_temp(data['personal_info']['foto'])

    # === TEMPLATE LOGIC ===
    
    # 1. MODERN SIDEBAR
    if settings['template_style'] == 'modern_sidebar':
        # Sidebar Background
        pdf.set_fill_color(r_prim, g_prim, b_prim)
        pdf.rect(0, 0, 65, 297, 'F')
        
        # --- SIDEBAR CONTENT ---
        pdf.set_text_color(255, 255, 255)
        
        # Photo
        curr_y = 20
        if temp_img_path:
            pdf.image(temp_img_path, x=12, y=20, w=40)
            curr_y += 45
        
        pdf.set_xy(5, curr_y)
        
        # Contact
        pdf.set_font(font, 'B', 14)
        pdf.cell(55, 10, "CONTACT", ln=True, align='L')
        pdf.set_font(font, '', 9)
        
        contacts = [
            data['personal_info']['email'],
            data['personal_info']['telepon'],
            data['personal_info']['alamat'],
            data['personal_info']['linkedin']
        ]
        for item in contacts:
            if item:
                pdf.set_x(5)
                pdf.multi_cell(55, 5, item)
                pdf.ln(2)
        
        pdf.ln(5)
        
        # Skills (Sidebar)
        if data['keahlian']:
            pdf.set_x(5)
            pdf.set_font(font, 'B', 14)
            pdf.cell(55, 10, "SKILLS", ln=True)
            pdf.set_font(font, '', 9)
            for skill in data['keahlian']:
                pdf.set_x(5)
                pdf.cell(55, 6, f"- {skill}", ln=True)

        # Languages (Sidebar)
        if data['bahasa']:
            pdf.ln(5)
            pdf.set_x(5)
            pdf.set_font(font, 'B', 14)
            pdf.cell(55, 10, "LANGUAGES", ln=True)
            pdf.set_font(font, '', 9)
            for lang in data['bahasa']:
                pdf.set_x(5)
                pdf.cell(55, 6, f"- {lang}", ln=True)

        # --- MAIN CONTENT ---
        pdf.set_xy(70, 20)
        pdf.set_text_color(r_prim, g_prim, b_prim)
        
        # Name
        pdf.set_font(font, 'B', 28)
        pdf.cell(0, 10, data['personal_info']['nama'].upper(), ln=True)
        
        # Position
        if data['personal_info']['posisi_target']:
            pdf.set_font(font, 'B', 16)
            pdf.set_text_color(r_sec, g_sec, b_sec)
            pdf.cell(0, 10, data['personal_info']['posisi_target'], ln=True)
        
        pdf.ln(5)
        
        # Summary
        if data['ringkasan']:
            pdf.set_text_color(0, 0, 0)
            pdf.set_font(font, '', 10)
            pdf.multi_cell(0, 5, data['ringkasan'])
            pdf.ln(10)
            
        # Helper to avoid repetitive code
        def add_section(title):
            pdf.set_x(70)
            pdf.set_font(font, 'B', 14)
            pdf.set_text_color(r_prim, g_prim, b_prim)
            pdf.cell(0, 8, title.upper(), border='B', ln=True)
            pdf.ln(2)
            
        # Experience
        if data['pengalaman']:
            add_section("Work Experience")
            for exp in data['pengalaman']:
                pdf.set_x(70)
                pdf.set_text_color(0, 0, 0)
                pdf.set_font(font, 'B', 11)
                pdf.cell(0, 6, exp.get('posisi', ''), ln=True)
                
                pdf.set_x(70)
                pdf.set_font(font, 'I', 10)
                pdf.set_text_color(100)
                pdf.cell(0, 5, f"{exp.get('perusahaan', '')} | {exp.get('periode', '')}", ln=True)
                
                pdf.set_x(70)
                pdf.set_font(font, '', 10)
                pdf.set_text_color(0)
                pdf.multi_cell(0, 5, exp.get('deskripsi', ''))
                pdf.ln(3)
        
        # Education
        if data['pendidikan']:
            pdf.ln(5)
            add_section("Education")
            for edu in data['pendidikan']:
                pdf.set_x(70)
                pdf.set_font(font, 'B', 11)
                pdf.set_text_color(0)
                pdf.cell(0, 6, edu.get('institusi', ''), ln=True)
                
                pdf.set_x(70)
                pdf.set_font(font, '', 10)
                pdf.cell(0, 5, f"{edu.get('gelar', '')} | {edu.get('tahun', '')}", ln=True)
                pdf.ln(2)

    # 2. CLASSIC / EXECUTIVE (Vertical Layout)
    else:
        # Header Center
        align = 'C' if settings['template_style'] == 'executive' else 'L'
        
        # Photo handling for vertical layout
        if temp_img_path and align == 'L':
             pdf.image(temp_img_path, x=170, y=10, w=30)
        
        pdf.set_font(font, 'B', 24)
        pdf.set_text_color(r_prim, g_prim, b_prim)
        pdf.cell(0, 10, data['personal_info']['nama'].upper(), ln=True, align=align)
        
        if data['personal_info']['posisi_target']:
            pdf.set_font(font, 'B', 14)
            pdf.set_text_color(r_sec, g_sec, b_sec)
            pdf.cell(0, 8, data['personal_info']['posisi_target'], ln=True, align=align)
            
        # Contact Bar
        pdf.set_font(font, '', 9)
        pdf.set_text_color(50)
        contacts = []
        if data['personal_info']['email']: contacts.append(data['personal_info']['email'])
        if data['personal_info']['telepon']: contacts.append(data['personal_info']['telepon'])
        if data['personal_info']['linkedin']: contacts.append("LinkedIn")
        if data['personal_info']['alamat']: contacts.append(data['personal_info']['alamat'])
        
        pdf.cell(0, 6, " | ".join(contacts), ln=True, align=align)
        pdf.line(10, pdf.get_y(), 200, pdf.get_y())
        pdf.ln(5)
        
        # Main Logic for Vertical
        def add_vertical_section(title):
            pdf.ln(3)
            pdf.set_font(font, 'B', 12)
            pdf.set_fill_color(r_prim, g_prim, b_prim)
            pdf.set_text_color(255)
            pdf.cell(0, 7, f"  {title.upper()}", ln=True, fill=True)
            pdf.ln(2)
            
        if data['ringkasan']:
            add_vertical_section("Professional Summary")
            pdf.set_text_color(0)
            pdf.set_font(font, '', 10)
            pdf.multi_cell(0, 5, data['ringkasan'])
            
        if data['pengalaman']:
            add_vertical_section("Work Experience")
            for exp in data['pengalaman']:
                pdf.set_text_color(0)
                pdf.set_font(font, 'B', 11)
                pdf.cell(140, 6, exp.get('posisi', ''))
                pdf.set_font(font, 'B', 10)
                pdf.cell(0, 6, exp.get('periode', ''), align='R', ln=True)
                
                pdf.set_font(font, 'I', 10)
                pdf.set_text_color(r_sec, g_sec, b_sec)
                pdf.cell(0, 5, exp.get('perusahaan', ''), ln=True)
                
                pdf.set_font(font, '', 10)
                pdf.set_text_color(0)
                pdf.multi_cell(0, 5, exp.get('deskripsi', ''))
                pdf.ln(3)

        if data['keahlian']:
            add_vertical_section("Skills")
            pdf.set_text_color(0)
            pdf.set_font(font, '', 10)
            pdf.multi_cell(0, 5, ", ".join(data['keahlian']))

        if data['pendidikan']:
            add_vertical_section("Education")
            for edu in data['pendidikan']:
                pdf.set_font(font, 'B', 11)
                pdf.set_text_color(0)
                pdf.cell(140, 6, edu.get('institusi', ''))
                pdf.set_font(font, 'B', 10)
                pdf.cell(0, 6, edu.get('tahun', ''), align='R', ln=True)
                
                pdf.set_font(font, '', 10)
                pdf.cell(0, 5, edu.get('gelar', ''), ln=True)
                pdf.ln(2)

    # Clean up temp file
    if temp_img_path and os.path.exists(temp_img_path):
        os.remove(temp_img_path)

    # Return buffer
    buffer = io.BytesIO()
    pdf.output(buffer)
    buffer.seek(0)
    return buffer

# --- WORD GENERATOR ---
def generate_word_doc(data, settings):
    doc = docx.Document()
    
    # Setup Colors
    r_prim, g_prim, b_prim = hex_to_rgb(settings['base_color'])
    r_sec, g_sec, b_sec = hex_to_rgb(settings['accent_color'])
    
    # Handle Photo (Temp File)
    temp_img_path = save_image_temp(data['personal_info']['foto'])

    # Style Choice
    if settings['template_style'] == 'modern_sidebar':
        # Table Layout
        table = doc.add_table(rows=1, cols=2)
        table.autofit = False
        table.columns[0].width = Inches(2.2)
        table.columns[1].width = Inches(4.8)
        
        # --- LEFT CELL (Sidebar) ---
        cell_l = table.cell(0, 0)
        
        # Add Photo if exists
        if temp_img_path:
            p = cell_l.add_paragraph()
            r = p.add_run()
            r.add_picture(temp_img_path, width=Inches(1.5))
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

        # Contact
        cell_l.add_paragraph("CONTACT").runs[0].bold = True
        
        contacts = [data['personal_info']['email'], data['personal_info']['telepon'], data['personal_info']['alamat']]
        for c in contacts:
            if c: cell_l.add_paragraph(c)
            
        # Skills
        if data['keahlian']:
            cell_l.add_paragraph("\nSKILLS").runs[0].bold = True
            for s in data['keahlian']:
                cell_l.add_paragraph(f"• {s}")

        # --- RIGHT CELL (Main) ---
        cell_r = table.cell(0, 1)
        
        # Name
        p = cell_r.add_paragraph()
        run = p.add_run(data['personal_info']['nama'].upper())
        run.bold = True
        run.font.size = Pt(24)
        run.font.color.rgb = RGBColor(r_prim, g_prim, b_prim)
        
        if data['personal_info']['posisi_target']:
            p = cell_r.add_paragraph(data['personal_info']['posisi_target'])
            p.runs[0].font.size = Pt(14)
            p.runs[0].font.color.rgb = RGBColor(r_sec, g_sec, b_sec)

        # Summary
        if data['ringkasan']:
            cell_r.add_paragraph("\nPROFESSIONAL SUMMARY").runs[0].bold = True
            cell_r.add_paragraph(data['ringkasan'])

        # Experience
        if data['pengalaman']:
            p = cell_r.add_paragraph("\nWORK EXPERIENCE")
            p.runs[0].bold = True
            p.runs[0].font.color.rgb = RGBColor(r_prim, g_prim, b_prim)
            
            for exp in data['pengalaman']:
                p = cell_r.add_paragraph()
                p.add_run(f"{exp.get('posisi','')}\n").bold = True
                p.add_run(f"{exp.get('perusahaan','')} | {exp.get('periode','')}\n").italic = True
                p.add_run(exp.get('deskripsi',''))

    else:
        # CLASSIC VERTICAL
        # Name
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(data['personal_info']['nama'].upper())
        run.bold = True
        run.font.size = Pt(22)
        run.font.color.rgb = RGBColor(r_prim, g_prim, b_prim)
        
        # Contact
        p = doc.add_paragraph(f"{data['personal_info']['email']} | {data['personal_info']['telepon']}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph().add_run().add_break() # Line Break

        # Sections
        def add_word_section(title, content_fn):
            p = doc.add_paragraph(title.upper())
            p.runs[0].bold = True
            p.runs[0].font.size = Pt(12)
            p.runs[0].font.color.rgb = RGBColor(r_prim, g_prim, b_prim)
            content_fn()
        
        if data['ringkasan']:
            add_word_section("Professional Summary", lambda: doc.add_paragraph(data['ringkasan']))
            
        if data['pengalaman']:
            def add_exp():
                for exp in data['pengalaman']:
                    p = doc.add_paragraph()
                    p.add_run(f"{exp.get('posisi','')}\n").bold = True
                    p.add_run(f"{exp.get('perusahaan','')} ({exp.get('periode','')})\n").italic = True
                    p.add_run(exp.get('deskripsi',''))
            add_word_section("Work Experience", add_exp)

        if data['keahlian']:
            add_word_section("Skills", lambda: doc.add_paragraph(", ".join(data['keahlian'])))

    # Cleanup
    if temp_img_path and os.path.exists(temp_img_path):
        os.remove(temp_img_path)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- HTML PREVIEW ---
def get_html_preview(data, settings):
    # CSS Variables
    base = settings['base_color']
    accent = settings['accent_color']
    
    # Photo Logic HTML
    photo_html = ""
    if data['personal_info']['foto']:
        photo_html = f'<img src="data:image/png;base64,{data["personal_info"]["foto"]}" style="width:100px; height:100px; border-radius:50%; object-fit:cover; margin-bottom:15px; border:3px solid white;">'
    
    if settings['template_style'] == 'modern_sidebar':
        return f"""
        <div style="display: flex; min-height: 800px; font-family: sans-serif; box-shadow: 0 0 15px rgba(0,0,0,0.1);">
            <div style="width: 30%; background-color: {base}; color: white; padding: 30px;">
                {photo_html}
                <div style="font-weight: bold; margin-bottom: 10px;">CONTACT</div>
                <div style="font-size: 0.9em; margin-bottom: 20px;">
                    <div>{data['personal_info']['email']}</div>
                    <div>{data['personal_info']['telepon']}</div>
                    <div>{data['personal_info']['alamat']}</div>
                </div>
                
                <div style="font-weight: bold; margin-bottom: 10px;">SKILLS</div>
                <div style="font-size: 0.9em;">
                    {''.join([f'<div style="margin-bottom:5px;">• {s}</div>' for s in data['keahlian']])}
                </div>
            </div>
            
            <div style="width: 70%; padding: 40px; background: white;">
                <h1 style="color: {base}; margin: 0;">{data['personal_info']['nama'].upper()}</h1>
                <h3 style="color: {accent}; margin-top: 5px;">{data['personal_info']['posisi_target']}</h3>
                
                <p style="margin-top: 20px; line-height: 1.6;">{data['ringkasan']}</p>
                
                <h3 style="color: {base}; border-bottom: 2px solid {base}; padding-bottom: 5px; margin-top: 30px;">WORK EXPERIENCE</h3>
                {''.join([f'''
                    <div style="margin-bottom: 15px;">
                        <div style="font-weight: bold; font-size: 1.1em;">{e.get('posisi','')}</div>
                        <div style="color: #666; font-style: italic; font-size: 0.9em;">{e.get('perusahaan','')} | {e.get('periode','')}</div>
                        <div style="margin-top: 5px;">{e.get('deskripsi','')}</div>
                    </div>
                ''' for e in data['pengalaman']])}
                
                <h3 style="color: {base}; border-bottom: 2px solid {base}; padding-bottom: 5px; margin-top: 30px;">EDUCATION</h3>
                 {''.join([f'''
                    <div style="margin-bottom: 10px;">
                        <div style="font-weight: bold;">{e.get('institusi','')}</div>
                        <div>{e.get('gelar','')} | {e.get('tahun','')}</div>
                    </div>
                ''' for e in data['pendidikan']])}
            </div>
        </div>
        """
    else:
        # CLASSIC PREVIEW
        return f"""
        <div style="padding: 50px; background: white; max-width: 800px; margin: 0 auto; box-shadow: 0 0 15px rgba(0,0,0,0.1); font-family: 'Times New Roman', serif;">
            <div style="text-align: center;">
                <h1 style="color: {base}; margin-bottom: 5px;">{data['personal_info']['nama'].upper()}</h1>
                <div style="margin-bottom: 10px;">{data['personal_info']['posisi_target']}</div>
                <div style="font-size: 0.9em; border-bottom: 1px solid #ccc; padding-bottom: 15px; margin-bottom: 20px;">
                    {data['personal_info']['email']} | {data['personal_info']['telepon']} | {data['personal_info']['alamat']}
                </div>
            </div>
            
            <h3 style="background: {base}; color: white; padding: 5px 10px; font-size: 1em; letter-spacing: 1px;">PROFESSIONAL SUMMARY</h3>
            <p>{data['ringkasan']}</p>
            
            <h3 style="background: {base}; color: white; padding: 5px 10px; font-size: 1em; letter-spacing: 1px; margin-top: 20px;">WORK EXPERIENCE</h3>
            {''.join([f'''
                <div style="margin-bottom: 15px;">
                    <div style="display:flex; justify-content:space-between; font-weight:bold;">
                        <span>{e.get('posisi','')}</span>
                        <span>{e.get('periode','')}</span>
                    </div>
                    <div style="font-style:italic; color: {accent}; margin-bottom: 5px;">{e.get('perusahaan','')}</div>
                    <div>{e.get('deskripsi','')}</div>
                </div>
            ''' for e in data['pengalaman']])}
        </div>
        """

# --- UI MAIN ---
st.title("🚀 CV Builder Pro Ultra")

# --- SIDEBAR (Progress & Settings) ---
with st.sidebar:
    st.header("📊 CV Health")
    score = calculate_ats_score(st.session_state.cv_data)
    st.progress(score/100)
    st.metric("ATS Score", f"{score}%")
    
    if score < 80:
        st.warning("⚠️ Tips: Add summary, experience, and at least 5 skills.")
    else:
        st.success("✅ Great Job! Your CV looks strong.")
        
    st.divider()
    if st.button("🗑️ Reset Data", type="primary"):
        st.session_state.cv_data = {k: (v if k != 'pengalaman' and k != 'pendidikan' and k != 'keahlian' else []) for k, v in st.session_state.cv_data.items()}
        st.rerun()

# --- MAIN TABS ---
tab_build, tab_design, tab_preview = st.tabs(["📝 Build Content", "🎨 Design & Style", "📥 Export"])

# 1. BUILD CONTENT
with tab_build:
    col1, col2 = st.columns([1, 2])
    with col1:
        st.subheader("Profile Photo")
        uploaded_file = st.file_uploader("Upload Image", type=['png', 'jpg', 'jpeg'])
        if uploaded_file:
            # Convert to base64 for session storage
            bytes_data = uploaded_file.getvalue()
            b64_str = base64.b64encode(bytes_data).decode()
            st.session_state.cv_data['personal_info']['foto'] = b64_str
            st.success("Photo Uploaded!")
        
        if st.session_state.cv_data['personal_info']['foto']:
             st.image(io.BytesIO(base64.b64decode(st.session_state.cv_data['personal_info']['foto'])), width=150)

    with col2:
        st.subheader("Personal Info")
        c1, c2 = st.columns(2)
        st.session_state.cv_data['personal_info']['nama'] = c1.text_input("Full Name", st.session_state.cv_data['personal_info']['nama'])
        st.session_state.cv_data['personal_info']['posisi_target'] = c2.text_input("Target Position", st.session_state.cv_data['personal_info']['posisi_target'])
        st.session_state.cv_data['personal_info']['email'] = c1.text_input("Email", st.session_state.cv_data['personal_info']['email'])
        st.session_state.cv_data['personal_info']['telepon'] = c2.text_input("Phone", st.session_state.cv_data['personal_info']['telepon'])
        st.session_state.cv_data['personal_info']['alamat'] = st.text_input("Address", st.session_state.cv_data['personal_info']['alamat'])

    st.subheader("Summary")
    st.session_state.cv_data['ringkasan'] = st.text_area("Professional Summary", st.session_state.cv_data['ringkasan'], height=100)

    st.subheader("Work Experience")
    if st.button("➕ Add Job"):
        st.session_state.cv_data['pengalaman'].append({})
    
    for i, exp in enumerate(st.session_state.cv_data['pengalaman']):
        with st.expander(f"Job #{i+1} - {exp.get('posisi', 'New Position')}", expanded=True):
            c1, c2 = st.columns(2)
            exp['posisi'] = c1.text_input(f"Position #{i}", exp.get('posisi', ''))
            exp['perusahaan'] = c2.text_input(f"Company #{i}", exp.get('perusahaan', ''))
            exp['periode'] = st.text_input(f"Period #{i}", exp.get('periode', ''))
            exp['deskripsi'] = st.text_area(f"Description #{i}", exp.get('deskripsi', ''))
            if st.button(f"Delete Job #{i}"):
                st.session_state.cv_data['pengalaman'].pop(i)
                st.rerun()

    st.subheader("Skills")
    skills_txt = st.text_area("Skills (comma separated)", ", ".join(st.session_state.cv_data['keahlian']))
    st.session_state.cv_data['keahlian'] = [s.strip() for s in skills_txt.split(",") if s.strip()]
    
    st.subheader("Education")
    if st.button("➕ Add Education"):
        st.session_state.cv_data['pendidikan'].append({})
    for i, edu in enumerate(st.session_state.cv_data['pendidikan']):
        with st.expander(f"Education #{i+1}", expanded=True):
            c1, c2 = st.columns(2)
            edu['institusi'] = c1.text_input(f"School/Uni #{i}", edu.get('institusi', ''))
            edu['gelar'] = c2.text_input(f"Degree #{i}", edu.get('gelar', ''))
            edu['tahun'] = st.text_input(f"Year #{i}", edu.get('tahun', ''))

# 2. DESIGN
with tab_design:
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Layout")
        layout = st.radio("Choose Layout", list(LAYOUTS.keys()), format_func=lambda x: LAYOUTS[x]['name'])
        st.session_state.settings['template_style'] = layout
        
        st.subheader("Fonts")
        font = st.selectbox("Font Family", list(FONTS.keys()))
        st.session_state.settings['font_family'] = font

    with c2:
        st.subheader("Theme Colors")
        base = st.color_picker("Primary Color", st.session_state.settings['base_color'])
        st.session_state.settings['base_color'] = base
        
        accent = st.color_picker("Accent Color", st.session_state.settings['accent_color'])
        st.session_state.settings['accent_color'] = accent

# 3. EXPORT
with tab_preview:
    st.subheader("Live Preview")
    html_prev = get_html_preview(st.session_state.cv_data, st.session_state.settings)
    st.components.v1.html(html_prev, height=800, scrolling=True)
    
    c1, c2 = st.columns(2)
    with c1:
        st.subheader("Download PDF")
        if st.button("📄 Generate PDF"):
            pdf_data = generate_pdf_enhanced(st.session_state.cv_data, st.session_state.settings)
            st.download_button("⬇️ Download PDF", pdf_data, file_name="my_cv.pdf", mime="application/pdf", type='primary')
            
    with c2:
        st.subheader("Download Word (Editable)")
        if st.button("📝 Generate Word"):
            try:
                docx_data = generate_word_doc(st.session_state.cv_data, st.session_state.settings)
                st.download_button("⬇️ Download Word", docx_data, file_name="my_cv.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", type='primary')
            except Exception as e:
                st.error(f"Error generating Word: {e}")
